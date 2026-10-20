"""Quality checks for the generated paper (run after build_docx.py).

Usage: python check_docx.py [path-to-docx]
Reports template leftovers (comments, placeholders, review notes), language, spacing and font problems.
"""
import re
import sys
import zipfile
from collections import Counter

import docx
from docx.oxml.ns import qn

PATH = sys.argv[1] if len(sys.argv) > 1 else "../docs/Transformers_in_NLP_Project_Paper.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
problems = []


def check(ok, msg):
    print(("  OK   " if ok else "  FAIL ") + msg)
    if not ok:
        problems.append(msg)


z = zipfile.ZipFile(PATH)
xml = z.read("word/document.xml").decode("utf8")
names = z.namelist()

print("Template leftovers")
check(not [n for n in names if "comment" in n.lower() or n.endswith("people.xml")], "no comment / people parts in the package")
for tag in ("commentRangeStart", "commentReference", "w:ins ", "w:del ", "w:highlight", "<w:sdt>", "w:footnoteReference"):
    check(xml.count(tag) == 0, f"no '{tag.strip()}' in document body")
for ph in ("Select study", "delete this note", "mano.vilniustech", "Student name", "TITLE OF THE FINAL", "Second-level", "Main text text",
           "Supervisor", "Consultant", "EXAMPLE", "Lorem", "TODO", "Right-click"):
    check(ph not in xml, f"no placeholder text '{ph}'")

d = docx.Document(PATH)
paras = [(p.style.name, p.text) for p in d.paragraphs]
text = "\n".join(t for _, t in paras)

print("Structure")
check(len(d.sections) == 2, f"2 sections (title page / contents and main text), found {len(d.sections)}")
check(not any(t.strip() == "Annotation" for _, t in paras), "no annotation page (course project)")
first_text = [t for _, t in paras[:30] if t.strip()]
check(sum("VILNIUS GEDIMINAS" in t for _, t in paras) == 1, "university name appears exactly once (single title page)")
check(any("Shashikanth" in t for t in first_text), "authors on the title page")

print("Language (English only)")
lt_chars = re.compile("[ąčęėįšųūžĄČĘĖĮŠŲŪŽ]")
body_lt = [t for s, t in paras if lt_chars.search(t) and s != "List Paragraph"]
# diacritics are allowed only inside author names of in-text citations
bad = [t for t in body_lt if not re.search(r"Kapočiūtė|Ulčar|Šikonja|Vileikytė|Lukoševičius|Stankevičius|Daniušis|Mulevičius|Nakvosas|Woźniak|Damaševičius|Lithuanian|Lithuania", t)]
check(not bad, "no Lithuanian text outside citations of Lithuanian authors" + (f" -> {bad[:2]}" if bad else ""))
check("UŽDUOTYSE" not in text and "TRANSFORMERIAI" not in text, "no Lithuanian title")

print("Spacing and typography")
dbl = [t[:60] for _, t in paras if "  " in t.strip()]
check(not dbl, f"no double spaces in paragraphs ({len(dbl)} found)" + (f" e.g. {dbl[:2]}" if dbl else ""))
sp = [t[:60] for _, t in paras if re.search(r" [,.;:)]", t) and "..." not in t]
check(not sp, f"no space before punctuation ({len(sp)} found)" + (f" e.g. {sp[:3]}" if sp else ""))
tabs = [t[:40] for s, t in paras if "\t" in t and not s.startswith(("BD_formule", "toc", "table of figures"))]
check(not tabs, f"no stray tab characters in body text ({len(tabs)} found)")
straight = text.count('"') + text.count("'")
check(straight == 0, f"typographic quotes only ({straight} straight quotes left)")
# consecutive empty paragraphs
run = max_run = 0
start = next(i for i, (_, t) in enumerate(paras) if t.strip() == "Contents")   # title page spacing is part of the template
for _, t in paras[start:]:
    run = run + 1 if not t.strip() else 0
    max_run = max(max_run, run)
check(max_run <= 3, f"no long runs of empty paragraphs (max {max_run})")

print("Fonts")
fonts, sizes = Counter(), Counter()
for p in d.paragraphs:
    for r in p.runs:
        if r.text.strip():
            fonts[r.font.name or p.style.font.name or "(style)"] += 1
            sizes[r.font.size.pt if r.font.size else (p.style.font.size.pt if p.style.font.size else None)] += 1
print("  fonts:", dict(fonts))
print("  sizes (None = inherits Normal 12 pt):", dict(sizes))
check(all(f in ("Times New Roman", "(style)", None) for f in fonts), "only Times New Roman in text runs")

print("Captions and references")
body_text = "\n".join(t for st, t in paras if not st.startswith(("toc", "table of figures")))
figs = len(re.findall(r"^Fig\. \d+\. ", body_text, re.M))
tabs_n = len(re.findall(r"^Table \d+\. ", body_text, re.M))
print(f"  figure captions: {figs}, table captions: {tabs_n}")
check("?" not in re.findall(r"(?:Table|Fig\.) (\S)", text) and "Table ?" not in text and "Fig. ?" not in text, "no unresolved cross-references")
refs = [t for s, t in paras if s == "List Paragraph"]
print(f"  reference entries: {len(refs)}")
cited = set()
for m in re.finditer(r"([A-ZČŠŽ][\w'’ščžėūųįąę\-]+)(?: et al\.| & [A-Z][\w\-ščžėūųįąęč]+)?,? \(?(\d{4})\)?", text):
    cited.add((m.group(1), m.group(2)))
missing = [(a, y) for a, y in sorted(cited) if not any(a in r and y in r for r in refs) and int(y) > 1990 and int(y) < 2026]
check(not missing, f"every in-text citation has a reference entry {missing[:8] if missing else ''}")

print()
print("RESULT:", "all checks passed" if not problems else f"{len(problems)} problem(s)")
sys.exit(1 if problems else 0)
